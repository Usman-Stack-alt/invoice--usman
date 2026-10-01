import io
import os
from pathlib import Path

os.environ.update(API_KEYS="test-key", ENV="test", GEMINI_API_KEY="", LLM_STATE_DIR="/tmp/invoice-llm-tests")

import pytest

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
def template_a() -> bytes:
    return (FIX / "template_a.jpg").read_bytes()


@pytest.fixture(scope="session")
def template_b() -> bytes:
    return (FIX / "template_b.jpg").read_bytes()


@pytest.fixture(scope="session")
def scanned_pdf(template_a) -> bytes:
    """Image-only PDF: forces the OCR path for PDFs."""
    from PIL import Image

    buf = io.BytesIO()
    Image.open(io.BytesIO(template_a)).convert("RGB").save(buf, "PDF", resolution=200)
    return buf.getvalue()


@pytest.fixture(scope="session")
def docx_invoice() -> bytes:
    """A born-digital Word invoice with a real table."""
    from docx import Document

    d = Document()
    d.add_heading("INVOICE", 0)
    d.add_paragraph("Invoice no: INV-1001")
    d.add_paragraph("Date of issue: 05/06/2021")
    d.add_paragraph("Due date: 06/05/2021")
    rows = [("Widget blue large", "2", "$10.00", "$20.00"), ("Gadget with long name that wraps", "3", "$10.00", "$30.00")]
    t = d.add_table(rows=1, cols=5)
    for c, h in zip(t.rows[0].cells, ("No.", "Description", "Quantity", "Price", "Total"), strict=False):
        c.text = h
    for i, r in enumerate(rows, 1):
        cells = t.add_row().cells
        for c, v in zip(cells, (str(i), *r), strict=False):
            c.text = v
    d.add_paragraph()
    for line in ("Subtotal $50.00", "Sales Tax 10% $5.00", "Total Due $55.00"):
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def from_to_pdf() -> bytes:
    """Digital PDF laid out like a real freelancer invoice: 'From (Seller)' / 'To (Buyer)' columns,
    key/value header, a table with a Total row, and a payment-details block."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    put = lambda x, y, t, size=10: page.insert_text((x, y), t, fontsize=size, fontname="helv")  # noqa: E731
    put(95, 60, "INVOICE", 12)
    put(95, 100, "From (Seller)")
    put(310, 100, "To (Buyer)")
    for i, t in enumerate(["Jane Example", "12 Example Road, Block 4", "Springfield, Exampleland", "Email:jane@example.com"]):
        put(95, 125 + 13 * i, t)
    for i, t in enumerate(["Acme Buyer GmbH", "Sample Gasse 1", "1010 Sampletown", "Austria", "VAT Nr: ATU00000000"]):
        put(310, 125 + 13 * i, t)
    for i, (k, v) in enumerate(
        [
            ("Invoice number:", "INV-2026-08"),
            ("Invoice date:", "2026-08-28"),
            ("Due date:", "2026-09-07"),
            ("Service period:", "22.07.2026-28.08.2026"),
            ("Currency", "EUR"),
        ]
    ):
        put(95, 215 + 25 * i, k)
        put(310, 215 + 25 * i, v)
    put(97, 360, "Description")
    put(450, 360, "Qty")
    put(490, 360, "Price (EUR)")
    put(97, 373, "IT Services - Development and implementation of an AI-based retrieval")
    put(450, 380, "1")
    put(490, 380, "625")
    put(97, 386, "system for the Company's platform")
    put(420, 400, "Total:")
    put(490, 400, "625")
    put(95, 440, "Payment details", 11)
    for i, (k, v) in enumerate(
        [
            ("Beneficiary", "Jane Example"),
            ("Bank", "Example Bank Ltd"),
            ("IBAN", "PK00EXMP0000000000000000"),
            ("BIC / SWIFT", "EXMPPKKA000"),
            ("Reference", "INV-2026-08"),
        ]
    ):
        put(95, 465 + 25 * i, k)
        put(310, 465 + 25 * i, v)
    return doc.tobytes()


@pytest.fixture(scope="session")
def receipt() -> bytes:
    """Thermal receipt photo: the rules flag it needs_review (a unit price had to be derived)."""
    return (FIX / "receipt.jpg").read_bytes()
